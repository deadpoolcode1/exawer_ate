"""Command-line entrypoint: `ate parse <file> [-o out.json]`."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from ate import __version__
from ate.errors import ATEParseError
from ate.parsers import parse


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(
        prog="ate",
        description="AI-Assisted Test Plan tool — M1: document parser",
    )
    p.add_argument("--version", action="version", version=f"ate {__version__}")
    sub = p.add_subparsers(dest="cmd", required=True)

    p_parse = sub.add_parser("parse", help="Parse a document into IR JSON")
    p_parse.add_argument("path", help="Path to PDF, DOCX, or TXT")
    p_parse.add_argument("-o", "--out", default=None,
                         help="Output JSON file (default: stdout)")
    p_parse.add_argument("--indent", type=int, default=2)
    p_parse.add_argument("--summary", action="store_true",
                         help="Print summary instead of full IR")

    p_plan = sub.add_parser("plan",
                            help="Generate a test plan xlsx from an input document (M1)")
    p_plan.add_argument("path", help="Path to input PDF/DOCX/TXT (single source)")
    p_plan.add_argument("-o", "--out", required=True,
                        help="Output xlsx path")
    p_plan.add_argument("--feature-name", default=None,
                        help="Override feature name (default: auto-detected)")
    p_plan.add_argument("--summary", action="store_true",
                        help="Print summary instead of writing xlsx")
    p_plan.add_argument("--no-ai", action="store_true",
                        help="Disable AI enrichment (force rule-based templates only)")
    p_plan.add_argument("--ai", action="store_true",
                        help="Force AI enrichment for any row not in ai_cache.json. "
                             "Routes through the backend chosen by --ai-backend.")
    p_plan.add_argument("--ai-backend", choices=("cli", "sdk"), default=None,
                        help="AI transport for enrichment. 'cli' (default) shells "
                             "out to `claude -p` using your local Claude Code "
                             "auth — no API key needed. 'sdk' uses the Anthropic "
                             "Python SDK and requires ANTHROPIC_API_KEY. Env var "
                             "ATE_AI_BACKEND overrides the default if this flag "
                             "is omitted.")
    p_plan.add_argument("--rfc", action="append", default=None, metavar="PATH",
                        help="Additional RFC source whose normative (MUST/SHALL) "
                             "clauses are extracted as requirements alongside the "
                             "spec anchors. Repeatable, e.g. "
                             "--rfc references/EVPN/rfc9785.txt --rfc references/EVPN/draft-...txt")
    p_plan.add_argument("--cli-doc", default=None, metavar="PATH",
                        help="EVPN CLI doc (DOCX). When provided, every config "
                             "command in the doc generates its own CLI Configuration "
                             "row family (happy-path / range / mutex / default / `no` "
                             "/ persistence / prerequisite). Replaces the generic CLI "
                             "templates and feeds the AI prompt with command evidence. "
                             "E.g. --cli-doc 'references/EVPN/EVPN CLI 1.00.docx'")

    p_pf = sub.add_parser("plan-feature",
                          help="Auto-discover SFS/CLI/RFCs under references/<NAME>/ "
                               "and generate the test plan xlsx for that feature")
    p_pf.add_argument("name", help="Feature folder name under references/ (e.g. EVPN)")
    p_pf.add_argument("-o", "--out", default=None,
                      help="Output xlsx path (default: plans/<NAME>_test_plan_with_RFCs.xlsx)")
    p_pf.add_argument("--root", default="references",
                      help="Root directory containing feature folders (default: references)")
    p_pf.add_argument("--feature-name", default=None,
                      help="Override the auto-detected feature display name")
    p_pf.add_argument("--summary", action="store_true",
                      help="Print summary instead of writing xlsx")
    p_pf.add_argument("--no-ai", action="store_true",
                      help="Disable AI enrichment (rule-based templates only)")
    p_pf.add_argument("--ai", action="store_true",
                      help="Force AI enrichment for rows not in ai_cache.json")
    p_pf.add_argument("--ai-backend", choices=("cli", "sdk"), default=None,
                      help="AI transport (cli = local Claude Code auth; sdk = ANTHROPIC_API_KEY)")
    p_pf.add_argument("--dry-run", action="store_true",
                      help="Print the resolved SFS/CLI/RFCs without running the planner")

    p_cg = sub.add_parser("codegen",
                          help="Generate the JSystem/Java EVPN test suite (M2) "
                               "from the SFS + CLI doc into cmp-tests-project "
                               "layout")
    p_cg.add_argument("-o", "--out", default="out/codegen",
                      help="Output root; files land under "
                           "<out>/cmp/tests/evpn/ (default: out/codegen)")
    p_cg.add_argument("--sfs",
                      default="references/EVPN/EVPN System Specification 1.00.docx",
                      help="EVPN System Functional Spec (.docx)")
    p_cg.add_argument("--cli-doc",
                      default="references/EVPN/EVPN CLI 1.00.docx",
                      help="EVPN CLI doc (.docx) — commands are grounded against it")
    p_cg.add_argument("--summary", action="store_true",
                      help="Print the plan without writing files")
    p_cg.add_argument("--selected-only", action="store_true",
                      help="Only emit tests the dirty queue marks SELECTED "
                           "(the SOW's 'code generation based on selected "
                           "tests'); refreshes and updates the queue")
    p_cg.add_argument("--queue", default=None,
                      help="Queue file path (default: out/codegen_queue.json)")
    p_cap = sub.add_parser("capture",
                           help="Run the suite's show commands on a real DUT "
                                "and record their output as expectations")
    p_cap.add_argument("--host", required=True, help="DUT management IP")
    p_cap.add_argument("--user", default="admin")
    p_cap.add_argument("--password", default="admin")
    p_cap.add_argument("--jump", default=None, metavar="USER@HOST",
                       help="SSH through this host (the lab is not routable "
                            "from a laptop), e.g. ilan@192.168.31.226")
    p_cap.add_argument("--out", default="out/captured_expectations.json")
    p_cap.add_argument("--sfs",
                       default="references/EVPN/EVPN System Specification 1.00.docx")
    p_cap.add_argument("--cli-doc",
                       default="references/EVPN/EVPN CLI 1.00.docx")

    p_cg.add_argument("--from-plan", default=None, metavar="XLSX",
                      help="Generated test plan to derive additional suites "
                           "from mechanically (no curated steps)")
    p_cg.add_argument("--plan-flows", default="", metavar="IDS",
                      help="Comma-separated flow IDs to derive from --from-plan, "
                           "e.g. FLOW-020,FLOW-021. Emitted as TCM<nnn>_ classes")

    p_q = sub.add_parser("queue",
                         help="Dirty queue — which generated tests are "
                              "selected, generated, approved or stale")
    p_q.add_argument("action",
                     choices=("status", "refresh", "select", "select-all",
                              "approve"),
                     help="status: show the table; refresh: recompute "
                          "fingerprints and mark stale; select/approve: "
                          "change state for the given test IDs")
    p_q.add_argument("test_ids", nargs="*",
                     help="Test IDs (flow IDs, e.g. FLOW-030)")
    p_q.add_argument("--queue", default=None, help="Queue file path")
    p_q.add_argument("--sfs",
                     default="references/EVPN/EVPN System Specification 1.00.docx")
    p_q.add_argument("--cli-doc",
                     default="references/EVPN/EVPN CLI 1.00.docx")
    p_q.add_argument("--rehash", action="store_true",
                     help="Digest source documents by content instead of "
                          "size+mtime (slower, catches touch-only rewrites)")

    p_m = sub.add_parser("match",
                         help="Run the pattern library over a generated test "
                              "plan and report how much of it maps to "
                              "executable steps")
    p_m.add_argument("xlsx", help="Generated test plan .xlsx")
    p_m.add_argument("--show-unmatched", type=int, default=0,
                     help="Print N unmatched rows (default 0)")

    args = p.parse_args(argv)

    if args.cmd == "parse":
        return _cmd_parse(args)
    if args.cmd == "plan":
        return _cmd_plan(args)
    if args.cmd == "plan-feature":
        return _cmd_plan_feature(args)
    if args.cmd == "codegen":
        return _cmd_codegen(args)
    if args.cmd == "queue":
        return _cmd_queue(args)
    if args.cmd == "match":
        return _cmd_match(args)
    if args.cmd == "capture":
        return _cmd_capture(args)
    return 2


def _cmd_capture(args) -> int:
    """Record real device output as expectations for the generated suite."""
    from ate.codegen.capture import capture_for_scripts  # noqa: PLC0415
    from ate.codegen.evpn_scripts import evpn_scripts  # noqa: PLC0415

    scripts = evpn_scripts()
    session = capture_for_scripts(scripts, args.host, args.user,
                                  args.password, jump=args.jump)
    print(f"host  : {session.host}")
    print(f"build : {session.build}")
    for c in session.results:
        mark = {"ok": "OK ", "empty": "EMPTY", "unsupported": "NO "}[c.status]
        print(f"  [{mark:5}] {c.expect_key:42} {c.command}")
        if c.note:
            print(f"           {c.note}")
    counts = session.by_status()
    print("\n" + "  ".join(f"{k}={v}" for k, v in sorted(counts.items())))
    usable = session.usable()
    print(f"usable expectations: {len(usable)} of {len(session.results)}")
    if not usable:
        print("\nNothing usable was captured, so no expectation was written. "
              "That is the correct outcome when the device does not implement "
              "the feature - freezing a rejection as an expectation would "
              "produce a test that passes by asserting the feature is absent.")
    print(f"saved: {session.save(args.out)}")
    return 0


def _queue_bits(args):
    """Shared queue plumbing: load the queue, refresh it against the current
    scripts, and return (queue, scripts, refresh report)."""
    from ate.codegen.evpn_scripts import evpn_scripts  # noqa: PLC0415
    from ate.codegen.queue import (  # noqa: PLC0415
        DEFAULT_QUEUE_PATH,
        Queue,
        digest_sources,
        digest_sources_full,
    )

    sources = [args.sfs, args.cli_doc]
    digest = (digest_sources_full(sources) if getattr(args, "rehash", False)
              else digest_sources(sources))
    q = Queue.load(args.queue or DEFAULT_QUEUE_PATH)
    scripts = evpn_scripts()
    report = q.refresh(scripts, digest)
    return q, scripts, report


def _cmd_queue(args) -> int:
    q, scripts, report = _queue_bits(args)

    if args.action == "select":
        if not args.test_ids:
            print("error: select needs at least one test ID")
            return 2
        unknown = q.select(args.test_ids)
        for u in unknown:
            print(f"warning: unknown test ID {u!r}")
    elif args.action == "select-all":
        picked = q.select_all()
        print(f"selected {len(picked)}: {', '.join(picked)}")
        print("(APPROVED tests were left alone; they hold reviewed work)")
    elif args.action == "approve":
        if not args.test_ids:
            print("error: approve needs at least one test ID")
            return 2
        unknown = q.approve(args.test_ids)
        for u in unknown:
            print(f"warning: unknown test ID {u!r}")

    for key, label in (("added", "new"), ("stale", "stale"),
                       ("removed", "gone from the catalog")):
        if report[key]:
            print(f"{label}: {', '.join(report[key])}")
    if report["stale_approved"]:
        print("WARNING: these were APPROVED and their inputs changed — "
              "regenerating will overwrite reviewed work: "
              f"{', '.join(report['stale_approved'])}")

    print()
    print(f"{'TEST':<12} {'STATE':<10} {'CLASS':<36} FINGERPRINT")
    for tid in sorted(q.entries):
        e = q.entries[tid]
        print(f"{e.test_id:<12} {e.state.value:<10} {e.class_name:<36} "
              f"{e.fingerprint}")
    counts = {k: v for k, v in q.summary().items() if v}
    print("\n" + "  ".join(f"{k}={v}" for k, v in counts.items()))
    print(f"queue: {q.save()}")
    return 0


def _cmd_match(args) -> int:
    import openpyxl  # noqa: PLC0415

    from ate.codegen.patterns import match_plan  # noqa: PLC0415

    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb["Test Plan Topics"]
    rows = []
    for r in list(ws.iter_rows(values_only=True))[1:]:
        action = (r[1] or "").strip()
        if not action:
            continue
        reqs = [x.strip() for x in (r[2] or "").split(",") if x.strip()]
        rows.append((action, reqs))

    rep = match_plan(rows)
    print(f"plan rows with an action : {len(rows)}")
    print(f"cross-reference rows     : {rep.skipped} (not executable; excluded)")
    print(f"automatable rows         : {rep.total}")
    print(f"mapped to typed steps    : {rep.matched}  ({rep.recall:.1%})")
    print()
    for kind, n in sorted(rep.by_kind().items(), key=lambda kv: -kv[1]):
        print(f"  {kind:<18} {n}")
    print(f"\nunmatched: {len(rep.unmatched())} (emitted as TODO stubs "
          "carrying the original sentence)")
    for u in rep.unmatched()[:args.show_unmatched]:
        print(f"  - {u.text[:110]}")
    return 0


def _cmd_codegen(args) -> int:
    """M2 code generation. Grounding failures are fatal, not warnings."""
    from ate.codegen import generate_evpn_suite  # noqa: PLC0415
    from ate.codegen.commands import UngroundedCommandError  # noqa: PLC0415

    try:
        plan_flows = [f.strip() for f in (args.plan_flows or "").split(",")
                      if f.strip()]
        if plan_flows and not args.from_plan:
            print("error: --plan-flows needs --from-plan <xlsx>")
            return 1
        result = generate_evpn_suite(args.sfs, args.cli_doc,
                                     plan_xlsx=args.from_plan,
                                     plan_flows=plan_flows)
    except UngroundedCommandError as e:
        print(f"error: {e}")
        return 1

    queue = None
    if args.selected_only:
        queue, _scripts, report = _queue_bits(args)
        if report["stale_approved"]:
            print("WARNING: overwriting reviewed work for "
                  f"{', '.join(report['stale_approved'])}")
        selected = set(queue.selected_ids())
        if not selected:
            print("nothing SELECTED in the queue — "
                  "run `ate queue select <TEST-ID>` first")
            queue.save()
            return 0
        keep = {s.class_name for s in result.scripts if s.flow_id in selected}
        result.scripts = [s for s in result.scripts if s.flow_id in selected]
        result.files = [f for f in result.files
                        if not f.class_name.startswith("TC")
                        or f.class_name in keep]
        print(f"generating only SELECTED tests: {', '.join(sorted(selected))}")

    print(f"scripts: {len(result.scripts)}   "
          f"steps: {result.n_steps}   files: {len(result.files)}")
    for script in result.scripts:
        todo = len(script.open_todos)
        print(f"  {script.flow_id}  {script.class_name}  "
              f"{len(script.steps)} steps"
              + (f"  ({todo} awaiting lab data)" if todo else ""))

    if result.warnings:
        print("\nCLI-doc anomalies carried into the generated code:")
        for w in result.warnings:
            print(f"  - {w}")

    if result.open_todos:
        print(f"\n{len(result.open_todos)} step(s) need real device output "
              "before their assertions mean anything:")
        for t in result.open_todos:
            print(f"  - {t}")

    if args.summary:
        return 0

    written = result.write(args.out)
    print(f"\nwrote {len(written)} file(s) under {args.out}/")
    for w in written:
        print(f"  {w}")

    if queue is not None:
        queue.mark_generated([s.flow_id for s in result.scripts])
        print(f"queue updated: {queue.save()}")
    return 0


def _cmd_parse(args) -> int:
    src = Path(args.path)
    try:
        doc = parse(src)
    except ATEParseError as e:
        print(f"error: {type(e).__name__}: {e}", file=sys.stderr)
        return 1

    if args.summary:
        print(f"path:        {doc.source_path}")
        print(f"format:      {doc.source_format}")
        print(f"schema:      {doc.schema_version}")
        print(f"blocks:      {len(doc.blocks)}")
        print(f"headings:    {len(doc.headings)}")
        print(f"paragraphs:  {len(doc.paragraphs)}")
        print(f"code blocks: {len(doc.code_blocks)}")
        print(f"tables:      {len(doc.tables)}")
        return 0

    payload = doc.model_dump(mode="json")
    text = json.dumps(payload, indent=args.indent, ensure_ascii=False)
    if args.out:
        Path(args.out).write_text(text + "\n", encoding="utf-8")
    else:
        print(text)
    return 0


def _print_rfc_crosscheck(plan) -> None:
    """Print the SFS-vs-ingested RFC gap warning (stderr) if a plan carries
    a cross-check result. Aleksey Burger, 2026-06-04."""
    from ate.planner.rfc_crosscheck import format_warning  # noqa: PLC0415
    catalog = plan.__dict__.get("_catalog")
    cc = getattr(catalog, "rfc_crosscheck", None) if catalog else None
    if cc is None:
        return
    msg = format_warning(cc)
    if msg:
        print(msg, file=sys.stderr)


def _print_cli_crosscheck(plan) -> None:
    """Report scrubbed ungrounded commands (stderr), and hard-alert if any
    survived. Ron/Yossi, 2026-06-24; scope confirmed by Ilan 2026-06-25."""
    from ate.planner.cli_crosscheck import (  # noqa: PLC0415
        format_removal_summary,
        format_warning,
    )
    removed = plan.__dict__.get("_cli_removed")
    if removed:
        msg = format_removal_summary(removed)
        if msg:
            print(msg, file=sys.stderr)
    cc = plan.__dict__.get("_cli_crosscheck")
    if cc is not None:
        warn = format_warning(cc)
        if warn:
            print(warn, file=sys.stderr)


def _cmd_plan(args) -> int:
    from ate.planner import generate_plan, generate_plan_to_xlsx
    src = Path(args.path)
    # use_ai: True = force API; False = rule-based; None = cache-only (default)
    if args.no_ai:
        use_ai: bool | None = False
    elif args.ai:
        use_ai = True
    else:
        use_ai = None  # cache-only by default
    rfc_paths = args.rfc if args.rfc else None
    cli_doc_path = args.cli_doc if args.cli_doc else None
    try:
        if args.summary:
            plan = generate_plan(src, feature_name=args.feature_name,
                                 use_ai=use_ai, rfc_paths=rfc_paths,
                                 cli_doc_path=cli_doc_path,
                                 ai_backend=args.ai_backend)
        else:
            plan = generate_plan_to_xlsx(src, args.out,
                                         feature_name=args.feature_name,
                                         use_ai=use_ai, rfc_paths=rfc_paths,
                                         cli_doc_path=cli_doc_path,
                                         ai_backend=args.ai_backend)
    except ATEParseError as e:
        print(f"error: {type(e).__name__}: {e}", file=sys.stderr)
        return 1

    _print_rfc_crosscheck(plan)
    _print_cli_crosscheck(plan)
    print(f"feature:        {plan.feature_name}")
    print(f"source:         {plan.source_path}")
    print(f"requirements:   {plan.n_requirements}")
    print(f"plan rows:      {plan.n_rows}")
    if not args.summary:
        print(f"xlsx written:   {args.out}")
    return 0


def discover_feature_inputs(folder: Path) -> dict:
    """Classify files under a feature folder into SFS / CLI doc / RFCs.

    Heuristics (case-insensitive on the basename):
      * `*[Tt]emplate*` xlsx → ignored (output templates).
      * Name starts with `rfc` or contains `draft-` → RFC.
        For RFCs we prefer the `.txt` form (the planner expects plain text).
        Sibling `.docx`/`.pdf` copies of the same RFC are silently dropped
        from the planner input (kept on disk for the parser parity tests).
      * Name contains "CLI" and ends in `.docx` → CLI doc.
      * Remaining `.docx` → SFS. Exactly one must remain.

    Returns: {"sfs": Path, "cli_doc": Path | None, "rfcs": list[Path]}.
    Raises FileNotFoundError / ValueError on misconfigured folders so the
    caller can print a helpful message.
    """
    if not folder.is_dir():
        raise FileNotFoundError(f"feature folder not found: {folder}")

    docx_files: list[Path] = []
    cli_docs: list[Path] = []
    rfc_by_stem: dict[str, dict[str, Path]] = {}

    for f in sorted(folder.iterdir()):
        if not f.is_file():
            continue
        name = f.name
        lname = name.lower()
        if lname.endswith(".xlsx") and "template" in lname:
            continue
        is_rfc = lname.startswith("rfc") or "draft-" in lname
        if is_rfc and f.suffix.lower() in {".txt", ".docx", ".pdf"}:
            rfc_by_stem.setdefault(f.stem, {})[f.suffix.lower()] = f
            continue
        if f.suffix.lower() == ".docx":
            if "cli" in lname:
                cli_docs.append(f)
            else:
                docx_files.append(f)

    # Prefer .txt for each RFC; fall back to .docx then .pdf.
    rfcs: list[Path] = []
    for stem, by_ext in sorted(rfc_by_stem.items()):
        rfcs.append(by_ext.get(".txt") or by_ext.get(".docx") or by_ext[".pdf"])

    if len(cli_docs) > 1:
        raise ValueError(
            f"multiple CLI docs found in {folder}: {[p.name for p in cli_docs]}. "
            f"Only one is supported per feature."
        )
    if not docx_files:
        raise ValueError(
            f"no SFS .docx found in {folder} "
            f"(expected exactly one non-CLI, non-RFC .docx)"
        )
    if len(docx_files) > 1:
        raise ValueError(
            f"multiple SFS candidates in {folder}: {[p.name for p in docx_files]}. "
            f"Rename the CLI doc to contain 'CLI', or move extra .docx files out."
        )

    return {
        "sfs": docx_files[0],
        "cli_doc": cli_docs[0] if cli_docs else None,
        "rfcs": rfcs,
    }


def _cmd_plan_feature(args) -> int:
    from ate.planner import generate_plan, generate_plan_to_xlsx

    folder = Path(args.root) / args.name
    try:
        inputs = discover_feature_inputs(folder)
    except (FileNotFoundError, ValueError) as e:
        print(f"error: {e}", file=sys.stderr)
        return 1

    sfs: Path = inputs["sfs"]
    cli_doc: Path | None = inputs["cli_doc"]
    rfcs: list[Path] = inputs["rfcs"]

    print(f"feature folder: {folder}")
    print(f"  SFS:        {sfs.name}")
    print(f"  CLI doc:    {cli_doc.name if cli_doc else '(none)'}")
    print(f"  RFCs:       {', '.join(p.name for p in rfcs) if rfcs else '(none)'}")

    if args.dry_run:
        # Surface the SFS-vs-ingested RFC gap without running the planner —
        # a lightweight parse of just the SFS is enough to reconcile.
        try:
            from ate.planner.rfc_crosscheck import (  # noqa: PLC0415
                format_warning,
                reconcile,
            )
            sfs_doc = parse(sfs)
            cc = reconcile(sfs_doc.full_text, [str(p) for p in rfcs])
            msg = format_warning(cc)
            if msg:
                print(msg, file=sys.stderr)
        except ATEParseError as e:
            print(f"warning: could not parse SFS for RFC cross-check: {e}",
                  file=sys.stderr)
        return 0

    out_path = args.out or f"plans/{args.name}_test_plan_with_RFCs.xlsx"
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)

    if args.no_ai:
        use_ai: bool | None = False
    elif args.ai:
        use_ai = True
    else:
        use_ai = None

    rfc_paths = [str(p) for p in rfcs] if rfcs else None
    cli_doc_path = str(cli_doc) if cli_doc else None

    try:
        if args.summary:
            plan = generate_plan(sfs, feature_name=args.feature_name,
                                 use_ai=use_ai, rfc_paths=rfc_paths,
                                 cli_doc_path=cli_doc_path,
                                 ai_backend=args.ai_backend)
        else:
            plan = generate_plan_to_xlsx(sfs, out_path,
                                         feature_name=args.feature_name,
                                         use_ai=use_ai, rfc_paths=rfc_paths,
                                         cli_doc_path=cli_doc_path,
                                         ai_backend=args.ai_backend)
    except ATEParseError as e:
        print(f"error: {type(e).__name__}: {e}", file=sys.stderr)
        return 1

    _print_rfc_crosscheck(plan)
    _print_cli_crosscheck(plan)
    print(f"feature:        {plan.feature_name}")
    print(f"requirements:   {plan.n_requirements}")
    print(f"plan rows:      {plan.n_rows}")
    if not args.summary:
        print(f"xlsx written:   {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
